
from pathlib import Path
import fitz
from docx import Document


def extract_resume_text(file_path):
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        text_parts = []

        with fitz.open(str(path)) as pdf:
            if pdf.is_encrypted:
                raise ValueError("This PDF is password-protected.")

            for page in pdf:
                text_parts.append(page.get_text())

        text = "\n".join(text_parts)

    elif extension == ".docx":
        document = Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(
                    cell.text for cell in row.cells
                )

        text = "\n".join(paragraphs)

    else:
        raise ValueError("Please upload a PDF or DOCX file.")

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text found. The PDF may be scanned or image-based."
        )

    return text
